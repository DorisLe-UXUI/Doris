#!/usr/bin/env python3
"""Builds KDL_SEO_Plan.xlsx — keyword map, page meta (read from dist/), technical checklist, schema map,
12-week content calendar, link-building targets, redirect map, open questions. Bilingual EN/VN headers."""
import os, re, json, html, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIST = os.path.join(ROOT, "dist")
OUT = os.path.join(ROOT, "KDL_SEO_Plan.xlsx")
NAVY, GOLD, CREAM, MAROON, INK = "0F2143", "F6B523", "F8F5EF", "5F0001", "1B2233"
thin = Side(style="thin", color="D9D2C3")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
wb.remove(wb.active)

def sheet(name, title, subtitle, headers, rows, widths, wrap_cols=()):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = Font(name="Inter", size=16, bold=True, color=NAVY)
    ws["A2"] = subtitle; ws["A2"].font = Font(name="Inter", size=10, italic=True, color="6B7285")
    hr = 4
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.font = Font(name="Inter", bold=True, color=GOLD, size=10)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    for r, row in enumerate(rows, hr + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name="Inter", size=10, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=(c in wrap_cols or c == len(headers)))
            cell.border = BORDER
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=CREAM)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[hr].height = 34
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    ws.sheet_view.showGridLines = False
    return ws

# ------------------------------------------------------------------ 00 README
readme_rows = [
    ["Mục đích · Purpose", "Kế hoạch SEO đi kèm bản redesign kingdomlandkids.com (3D hero · UI/UX · SEO). Mỗi sheet là một lớp việc: từ khóa → meta từng trang → kỹ thuật → schema → nội dung → backlink → redirect → câu hỏi mở.", "SEO plan shipped with the kingdomlandkids.com redesign. Each sheet is one layer: keywords → per-page meta → technical → schema → content → links → redirects → open questions."],
    ["01_Keyword_Map", "Cụm từ khóa theo funnel, intent và trang đích. Volume/KD điền từ GSC + Ahrefs/Semrush sau khi có access — cột Priority là đánh giá định tính.", "Keyword clusters by funnel stage, intent and target page. Fill volume/KD from GSC + Ahrefs/Semrush once access exists — Priority is a qualitative call."],
    ["02_Page_Meta", "Đọc trực tiếp từ dist/ sau khi build: title, meta description, H1, canonical, số ký tự, schema. Đây là 'nguồn sự thật' để QA.", "Read straight from the built dist/: title, description, H1, canonical, character counts, schema. Use as the QA source of truth."],
    ["03_Technical", "Checklist kỹ thuật: cái gì đã có trong build, cái gì phải làm lúc launch (DNS, GSC, analytics, redirect).", "Technical checklist: what the build already does vs. what must happen at launch (DNS, GSC, analytics, redirects)."],
    ["04_Schema_Map", "Loại structured data trên từng trang và cách kiểm tra.", "Structured-data types per page and how to validate."],
    ["05_Content_Calendar", "12 tuần blog (Sep 14 → Dec 6, 2026) bám mùa: back-to-school, Thanksgiving, Advent/Christmas, chuẩn bị MLK Day & Black History Month.", "12-week blog calendar (Sep 14 → Dec 6, 2026) following the season: back-to-school, Thanksgiving, Advent/Christmas, MLK Day & Black History Month prep."],
    ["06_Links_Distribution", "Nơi kiếm backlink/đề cập chất lượng: roundup blogs Christian mom, directory faith apps, church partners, YouTube/IG bio.", "Where quality links/mentions come from: Christian-mom roundup blogs, faith-app directories, church partners, YouTube/IG bios."],
    ["07_Redirect_Map", "Khung map URL cũ → mới. Cần crawl site hiện tại (Screaming Frog) để điền cột URL cũ.", "Old → new URL map. Crawl the current site (Screaming Frog) to fill the old-URL column."],
    ["08_Open_Questions", "Những gì cần Doris/Nick xác nhận trước khi go-live.", "What Doris/Nick must confirm before go-live."],
    ["Cách dùng · How to use", "1) Duyệt 08 trước → chốt giá, tuổi, URL. 2) Dev chạy 03 lúc launch. 3) Content team chạy 05 hàng tuần. 4) Sau 4 tuần: đối chiếu GSC với 01 và cập nhật Priority.", "1) Clear sheet 08 first → lock pricing, ages, URLs. 2) Dev runs 03 at launch. 3) Content runs 05 weekly. 4) After 4 weeks: compare GSC with 01 and re-rank Priority."],
]
sheet("00_README", "KDL SEO PLAN · kingdomlandkids.com redesign", f"Generated {datetime.date.today().isoformat()} · SentryX / Doris Le · Sheets are bilingual VN · EN",
      ["Sheet / Mục", "Tiếng Việt", "English"], readme_rows, [24, 70, 70], wrap_cols=(2, 3))

# ------------------------------------------------------------------ 01 Keyword map
K = []  # cluster, keyword, intent, funnel, priority, target page, notes
def kw(cluster, words, intent, funnel, prio, page, note=""):
    for w in words:
        K.append([cluster, w, intent, funnel, prio, page, note])
kw("Brand", ["kingdomland kids", "kingdom land kids", "kingdomland kids app", "kingdomland kids login", "kingdomland kids free", "kingdomland kids marvelous light", "kingdomland kids youtube"], "Navigational", "Bottom", "P0", "/ , go.kingdomlandkids.com", "Own every brand variant; add 'Kingdomland Kids' to every title tag suffix (done).")
kw("Category", ["christian streaming for kids", "christian streaming service for kids", "christian kids streaming app", "faith based streaming for kids", "christian netflix for kids", "christian alternative to youtube kids", "ad free kids streaming christian"], "Commercial", "Middle", "P0", "/", "Home H1 kicker + description target this cluster.")
kw("Shows", ["christian shows for kids", "christian cartoons for kids", "bible cartoons for kids", "christian tv shows for toddlers", "christian shows for preschoolers", "animated bible stories for kids", "bible stories for kids videos"], "Commercial / Informational", "Middle", "P0", "/shows/ , blog checklist", "Shows index + '6-point checklist' article. Add a '10 best…' roundup only if we are comfortable naming others.")
kw("Marvelous Light", ["marvelous light kids show", "marvelous light lucy and leo", "fruits of the spirit cartoon", "bible cartoon talking animals"], "Navigational / Informational", "Middle", "P1", "/shows/marvelous-light/", "Long-tail is low volume but converts; TVSeries schema in place.")
kw("Worship & verses", ["kids worship songs", "bible memory verse songs for kids", "scripture songs for kids", "memory verse songs preschool", "christian songs for toddlers"], "Informational", "Top", "P1", "/shows/worship-and-memory-verses/", "Consider a free sampler playlist on YouTube linking back.")
kw("Fruits of the Spirit", ["fruits of the spirit for kids", "fruit of the spirit lesson for kids", "fruits of the spirit explained for children", "galatians 5:22 for kids"], "Informational", "Top", "P1", "/blog/fruits-of-the-spirit-for-kids/ → /shows/fruits-of-virtue/", "Evergreen; Sunday-school teachers search this year-round.")
kw("Family devotional", ["family devotional for preschoolers", "5 minute family devotional", "how to do family devotions with toddlers", "family bible time with young kids", "bedtime bible routine kids"], "Informational", "Top", "P1", "/blog/five-minute-family-devotional/", "Highest-intent parent pain point in the Q3 plan ('faith-filled screen time').")
kw("Black History Beats", ["black history songs for kids", "jackie robinson song for kids", "bessie coleman for kids", "mlk songs for kids", "black history month videos for kids christian", "juneteenth videos for kids"], "Informational / Seasonal", "Top", "P1", "/shows/black-history-beats/", "Seasonal peaks: Jan (MLK), Feb (BHM), Jun (Juneteenth). Publish 6 weeks ahead.")
kw("Audience-specific", ["christian shows for black kids", "bible cartoons black characters", "african american christian kids videos", "black christian family streaming"], "Commercial", "Middle", "P0", "/ , /about/", "Under-served, low competition, exact brand fit. About page carries this story.")
kw("Church", ["children's ministry videos", "sunday school videos for kids", "kids church videos free", "children's church curriculum videos", "church streaming for kids families"], "Commercial (B2B)", "Middle", "P1", "/churches/", "Pair with outreach to children's pastors; partner codes already in the model.")
kw("Comparison", ["minno vs yippee", "yippee tv alternative", "minno alternative", "best christian streaming service for kids 2026", "pure flix kids alternative"], "Commercial", "Middle", "P2", "future blog", "Only with verified facts about competitors — the May plan dropped unverified competitor pricing. Keep to features, not prices.")
kw("Preschool", ["christian abc videos", "christian counting songs toddlers", "safe preschool videos no ads", "christian preschool learning videos"], "Informational", "Top", "P2", "/shows/abcs-and-numbers/", "Broad and competitive; win on 'safe/no ads' angle.")
sheet("01_Keyword_Map", "Keyword map · Bản đồ từ khóa", "Priority: P0 = ship & measure now · P1 = next 90 days · P2 = opportunistic. Fill Volume / KD from GSC + Ahrefs once access exists.",
      ["Cluster / Cụm", "Keyword / Từ khóa", "Intent / Ý định", "Funnel", "Priority", "Target page / Trang đích", "Notes / Ghi chú", "Volume (fill)", "KD (fill)"],
      [r + ["", ""] for r in K], [18, 40, 22, 10, 9, 36, 60, 12, 10], wrap_cols=(6, 7))

# ------------------------------------------------------------------ 02 Page meta (from dist)
rows = []
for dirpath, _, files in os.walk(DIST):
    for fn in files:
        if fn != "index.html" and fn != "404.html":
            continue
        p = os.path.join(dirpath, fn)
        t = open(p, encoding="utf-8").read()
        rel = os.path.relpath(p, DIST).replace(os.sep, "/")
        url = "/" if rel == "index.html" else ("/" + rel.replace("/index.html", "/") if rel.endswith("index.html") else "/" + rel)
        g = lambda rx: (re.search(rx, t, re.S) or [None, ""])[1]
        title = html.unescape(g(r"<title>(.*?)</title>"))
        desc = html.unescape(g(r'<meta name="description" content="(.*?)"'))
        h1 = html.unescape(re.sub(r"<[^>]+>", " ", g(r"<h1[^>]*>(.*?)</h1>")))
        h1 = re.sub(r"\s+", " ", h1).strip()
        canon = g(r'<link rel="canonical" href="(.*?)"')
        robots = g(r'<meta name="robots" content="(.*?)"')
        ld = g(r'<script type="application/ld\+json">(.*?)</script>')
        try:
            types = ", ".join(x["@type"] if isinstance(x["@type"], str) else "/".join(x["@type"]) for x in json.loads(ld)["@graph"])
        except Exception:
            types = "INVALID"
        h2 = len(re.findall(r"<h2", t))
        words = len(re.sub(r"<[^>]+>", " ", re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", t, flags=re.S)).split())
        flag = []
        if len(title) > 60: flag.append(f"title {len(title)} chars (>60)")
        if len(desc) > 160: flag.append(f"description {len(desc)} chars (>160)")
        if not h1: flag.append("no H1")
        rows.append([url, title, len(title), desc, len(desc), h1, h2, words, canon, robots, types, "; ".join(flag) or "OK"])
rows.sort(key=lambda r: (r[0].count("/"), r[0]))
sheet("02_Page_Meta", "Per-page meta · Meta từng trang (đọc từ dist/)", "Google shows ~60 chars of a title and ~155–160 of a description; longer is not an error, it just gets truncated. Series titles run long on purpose (brand suffix) — trim if CTR is weak.",
      ["URL", "Title tag", "Title chars", "Meta description", "Desc chars", "H1", "H2 count", "Words (approx)", "Canonical", "Robots", "JSON-LD types", "QA"],
      rows, [34, 60, 8, 70, 8, 44, 7, 9, 44, 22, 46, 30], wrap_cols=(2, 4, 6, 11, 12))

# ------------------------------------------------------------------ 03 Technical
T = [
 ["Crawl & index", "Sitemap.xml auto-generated from the build, lastmod = build date", "Done in build", "Dev", "Submit in Google Search Console + Bing Webmaster after launch."],
 ["Crawl & index", "robots.txt allows all, blocks /privacy/ & /terms/ placeholders, points to sitemap", "Done in build", "Dev", "Remove the Disallow lines once real legal pages exist."],
 ["Crawl & index", "Canonical + hreflang (en, x-default) on every page", "Done in build", "—", ""],
 ["Crawl & index", "Custom 404.html (noindex)", "Done in build", "Dev", "Netlify/GH Pages pick up 404.html automatically; Vercel/Cloudflare need a rewrite rule."],
 ["Crawl & index", "301 redirects from every old URL (see 07_Redirect_Map)", "To do at launch", "Dev", "Crawl current site first. Never launch without this — it protects existing rankings."],
 ["Crawl & index", "Force https + one host (www → apex or apex → www) + trailing-slash consistency", "To do at launch", "Dev / DNS", "Build uses trailing slashes (/pricing/). Configure the host to redirect /pricing → /pricing/."],
 ["Crawl & index", "Google Search Console: DNS-verify kingdomlandkids.com (covers go. subdomain too)", "To do at launch", "Nick / Dev", "Also verify go.kingdomlandkids.com as a URL-prefix property to see funnel pages."],
 ["Structured data", "JSON-LD @graph: Organization, WebSite, WebApplication+Offers, TVSeries, FAQPage, Article, BreadcrumbList, AboutPage, ContactPage", "Done in build", "—", "Validate with https://validator.schema.org and Google Rich Results Test after launch."],
 ["Structured data", "Organization logo 512×512 PNG, sameAs → YouTube/IG/FB", "Done (placeholder logo)", "Doris", "Swap logo-512.png for the real wordmark-on-navy asset; confirm IG/FB URLs."],
 ["On-page", "One H1 per page with the primary keyword; kicker keeps SEO phrase inside the H1 on the home page", "Done in build", "—", ""],
 ["On-page", "Title ≤ 60 chars where possible; brand suffix on every page; unique meta descriptions", "Done (see 02 QA column)", "Doris", "Series titles intentionally long; trim if GSC shows truncation hurting CTR."],
 ["On-page", "Internal linking: every show ↔ related shows, blog ↔ shows, pricing ↔ churches, footer hub", "Done in build", "—", ""],
 ["On-page", "Image alt / aria on decorative SVG (aria-hidden) and meaningful controls", "Done in build", "—", "When real key art replaces the emblems, give each <img> a descriptive alt (series name + scene)."],
 ["Performance", "Hero is code (WebGL), no LCP image; fonts preconnected with display=swap; single CSS; deferred JS", "Done in build", "—", "Measure with PageSpeed Insights after deploy. Target: LCP < 2.5s mobile, CLS < 0.1, INP < 200ms."],
 ["Performance", "Three.js only loads on the home page; DPR capped; render pauses off-screen; reduced-motion = static frame", "Done in build", "—", "Vendored (687 KB min, ~170 KB brotli). Served self-hosted, cacheable for a year."],
 ["Performance", "Enable Brotli/gzip + long cache headers for /assets/ (immutable) on the host", "To do at launch", "Dev", "Netlify/Vercel/Cloudflare do this by default."],
 ["Performance", "Convert real show thumbnails to AVIF/WebP ≤ 60 KB, 16:10, with width/height attributes", "To do (when art arrives)", "Doris / Dev", "Prevents CLS; keeps mobile LCP fast."],
 ["Analytics", "GA4 property + Meta Pixel on the site (partials/analytics.html)", "To do at launch", "Dev", "Placeholders only — no fake IDs shipped."],
 ["Analytics", "App must fire Lead (free signup), Subscribe / StartTrial, Purchase (value, currency) to Meta + GA4", "To do — blocking ads efficiency", "Khai / Dev", "Sep 7 audit: pixel only sends PageView + AddToCart. Meta cannot optimise for subscribers until this ships."],
 ["Analytics", "UTM discipline: site CTAs → go.kingdomlandkids.com with utm_source=site&utm_medium=cta&utm_campaign=<page>", "Partly done (footer form)", "Doris", "Add UTMs to every CTA href in site.json if the app can read them."],
 ["Security", "HSTS, X-Content-Type-Options, Referrer-Policy, CSP (allow fonts.googleapis/gstatic + self)", "To do at launch", "Dev", "Static host header rules."],
 ["Accessibility", "Skip link, focus-visible styles, aria labels, colour contrast on navy/cream/gold (AA)", "Done in build", "—", "Gold-on-navy small text kept ≥ 700 weight for contrast."],
 ["Content", "Blog: 3 seed articles live; 12-week calendar in sheet 05", "Done / ongoing", "Doris", "One article per week beats twelve in a day."],
 ["Content", "Add a VideoObject for each episode once trailers/previews are embeddable", "Later", "Dev", "Big opportunity: video rich results for '<story> bible story for kids'."],
 ["Local / entity", "Google Business Profile not needed (no storefront); focus on Knowledge Panel via Organization schema + Wikidata/Crunchbase entries", "Later", "Nick", ""],
]
sheet("03_Technical", "Technical checklist · Checklist kỹ thuật", "Status: Done in build · To do at launch · Later. Owner column is a suggestion.",
      ["Area / Nhóm", "Item / Hạng mục", "Status", "Owner", "Notes / Ghi chú"], T, [16, 80, 22, 14, 70], wrap_cols=(2, 5))

# ------------------------------------------------------------------ 04 Schema map
S = [
 ["All pages", "Organization (@id #organization), WebSite", "name, url, logo, email, sameAs, contactPoint", "Rich Results Test → 'Organization' has no rich result but feeds the Knowledge Panel."],
 ["/ and /pricing/", "WebApplication + SoftwareApplication with 3 × Offer", "applicationCategory, operatingSystem, offers[price, priceCurrency, UnitPriceSpecification], isFamilyFriendly", "Software app rich result eligible; keep prices in sync with site.json."],
 ["/shows/<slug>/", "TVSeries + BreadcrumbList", "name, description, numberOfEpisodes, typicalAgeRange, genre, provider", "Add Season / TVEpisode + VideoObject later per episode."],
 ["/faq/", "FAQPage + BreadcrumbList", "Question / Answer pairs (9)", "FAQ rich results are limited to authoritative sites since 2023 — still useful for AI overviews."],
 ["/blog/<slug>/", "Article + BreadcrumbList", "headline, datePublished, author (Organization), image", "Add a Person author later if the team wants bylines."],
 ["/about/", "AboutPage + BreadcrumbList", "url, name", ""],
 ["/contact/", "ContactPage + BreadcrumbList", "url, name", ""],
 ["/churches/", "BreadcrumbList (consider Service later)", "—", "If church pricing becomes public, add an Offer with eligibleCustomerType."],
]
sheet("04_Schema_Map", "Structured data map · Bản đồ schema", "Everything is emitted as one JSON-LD @graph per page by build.py.",
      ["Page(s)", "Types", "Key properties", "Notes"], S, [22, 46, 70, 70], wrap_cols=(2, 3, 4))

# ------------------------------------------------------------------ 05 Content calendar
start = datetime.date(2026, 9, 14)
C = [
 ("Back-to-school mornings: a 10-minute faith routine before the bus", "morning routine christian family kids", "/shows/worship-and-memory-verses/", "Buổi sáng đi học: routine 10 phút với worship + memory verse"),
 ("How to explain the Exodus to a 6-year-old (without the scary parts)", "exodus story for kids explained", "/shows/bible-adventures/", "Giải thích Xuất Hành cho trẻ 6 tuổi"),
 ("Screen-time rules that actually work for Christian families", "christian family screen time rules", "/ (why parents choose)", "Quy tắc screen time cho gia đình Cơ Đốc"),
 ("Psalm 23 for kids: a bedtime walk-through", "psalm 23 for kids", "/shows/bible-adventures/", "Thi Thiên 23 cho trẻ trước giờ ngủ"),
 ("What Black History Beats teaches that textbooks skip", "black history for kids christian", "/shows/black-history-beats/", "Black History Beats dạy gì mà sách giáo khoa bỏ qua"),
 ("A thankful jar: a 4-week Thanksgiving devotional for little kids", "thanksgiving devotional for kids", "/blog/five-minute-family-devotional/", "Lọ tạ ơn: devotional 4 tuần mùa Thanksgiving"),
 ("Is the annual plan worth it? The honest math for one family", "kingdomland kids pricing annual", "/pricing/", "Gói năm có đáng không: phép tính thật"),
 ("Advent with a 4-year-old: 24 tiny moments, no crafts required", "advent devotional for toddlers", "/shows/marvelous-light/ (Christmas)", "Mùa Vọng với bé 4 tuổi: 24 khoảnh khắc nhỏ"),
 ("The Christmas story, told by Lucy and Leo (episode guide)", "christmas bible story cartoon for kids", "/shows/marvelous-light/", "Câu chuyện Giáng Sinh qua Lucy & Leo"),
 ("Gift idea: a year of Bible adventures + a storybook at the door", "christian gift for kids subscription", "/pricing/", "Ý tưởng quà: gói năm + sách"),
 ("New Year, one verse: picking a family verse for 2027", "family bible verse for the year", "/shows/worship-and-memory-verses/", "Câu Kinh Thánh của gia đình cho 2027"),
 ("Getting ready for MLK Day: songs, stories and one hard question", "mlk day activities for kids christian", "/shows/black-history-beats/", "Chuẩn bị MLK Day: bài hát, câu chuyện, một câu hỏi khó"),
]
rows = []
for i, (title, kwd, link, vn) in enumerate(C):
    d = start + datetime.timedelta(weeks=i)
    rows.append([f"W{i+1}", d.isoformat(), title, vn, kwd, link, "Try free — no card → go.kingdomlandkids.com", "Blog post (700–900 words) + 1 IG carousel + 1 Reel hook + Klaviyo note"])
sheet("05_Content_Calendar", "12-week content calendar · Lịch nội dung 12 tuần", "One post per week, published Monday. Every post links to one show page and the free plan. Repurpose into the Q3/Q4 social calendar.",
      ["Week", "Publish date", "Title (EN)", "Góc nội dung (VN)", "Target keyword", "Internal link", "CTA", "Repurpose"], rows, [7, 13, 58, 46, 34, 36, 36, 40], wrap_cols=(3, 4, 8))

# ------------------------------------------------------------------ 06 Links
L = [
 ["Roundup blogs (Christian moms)", "raisingchristiankids.com · equippinggodlywomen.com · mindyjonesblog.com · themillennialsahm.com · proverbialhomemaker.com · crosswalk.com", "Pitch 'Kingdomland Kids' for their 'Christian streaming / shows for kids' lists (they already rank; Yippee & Minno are on them).", "P0", "Doris"],
 ["Directories", "faith.tools/for-kids · Christian app directories · Common Sense Media listing", "Submit listing with logo, screenshots, pricing.", "P0", "Doris"],
 ["Church partners", "Every partner church website / newsletter", "Ask for a 'Recommended for families' link to /churches/ or the code page.", "P1", "Nick"],
 ["Owned channels", "YouTube (159K) channel + video descriptions · IG bio · FB page", "Link to kingdomlandkids.com (not only go.), so the root domain accrues authority.", "P0", "Jojo / Linh"],
 ["Creators / ambassadors", "Ambassador testimonial creators, 'next Miss Rachel' creator", "Ask for a link in bio + post copy pointing to a show page.", "P1", "Matt / Nick"],
 ["Press", "Christian parenting podcasts, Black family media, local TX/CA/FL/NY faith outlets", "Angle: 'the ad-free Bible app built by and for Black Christian families'.", "P2", "Nick"],
 ["Partnerships", "CoinCept (partnership videos already in v1.4), PlayWatchKids sister channel", "Cross-link footers and about pages.", "P1", "Nick"],
]
sheet("06_Links_Distribution", "Links & distribution · Backlink và phân phối", "Quality over quantity: ten relevant Christian-parenting links beat a thousand directory spam links.",
      ["Type", "Targets", "Play", "Priority", "Owner"], L, [26, 70, 70, 9, 14], wrap_cols=(2, 3))

# ------------------------------------------------------------------ 07 Redirects
R = [
 ["https://kingdomlandkids.com/", "/", "Home", "Keep"],
 ["(crawl current site)", "/shows/", "Any 'shows', 'videos', 'library' page", "301"],
 ["(crawl current site)", "/pricing/", "Any 'plans', 'subscribe', 'pricing' page", "301"],
 ["(crawl current site)", "/about/", "About / our story", "301"],
 ["(crawl current site)", "/faq/", "FAQ / help", "301"],
 ["(crawl current site)", "/contact/", "Contact", "301"],
 ["go.kingdomlandkids.com/welcome", "keep (app)", "Paid-ad landing page stays on the app subdomain", "No change — but add canonical + noindex if it duplicates /pricing/"],
 ["Login / account URLs", "keep (app)", "Header 'Log in' currently points to go.kingdomlandkids.com — replace with the real login URL", "Update site.json › urls.login"],
]
sheet("07_Redirect_Map", "Redirect map · Bản đồ redirect", "Fill the old-URL column from a Screaming Frog crawl of the live site before launch. Every indexed old URL needs a 301.",
      ["Old URL", "New URL", "What it was", "Action"], R, [44, 22, 60, 50], wrap_cols=(3, 4))

# ------------------------------------------------------------------ 08 Open questions
Q = [
 ["Pricing", "Q3 plan says $9.99/mo ($7.99 first month) + $69.99/yr. Slack (Jun 14) says the live site is $7.99/mo + $69.99/yr. Which is live today?", "site.json › pricing", "Nick"],
 ["Ages", "Site says 'ages 2–10' overall, 4–8 for most shows. OK to publish?", "site.json › description, hero kicker, FAQ", "Doris"],
 ["Devices", "Copy says 'browser on phone/tablet/laptop + TV via casting'. Are there native apps (iOS/Android/Roku/Fire TV)? If yes we add store badges + SoftwareApplication per platform.", "Home devices section, FAQ", "Khai"],
 ["Login URL", "Real login URL for the header button (currently go.kingdomlandkids.com).", "site.json › urls.login", "Khai"],
 ["Social URLs", "Confirm Instagram and Facebook handles (YouTube confirmed).", "site.json › urls", "Jojo"],
 ["Testimonials", "Two quotes carried over from the current site — attribution / permission to keep using them?", "site.json › testimonials", "Nick"],
 ["Series facts", "Bible Adventures '15 episodes', Little Agent '8 missions', ABCs & Numbers count, Bible in Clay status (new? live?), Grandpa Moses series — confirm counts and names.", "site.json › series", "Matt"],
 ["Storybook", "Annual bonus ships to U.S. addresses only? Any timeline?", "Pricing page note, FAQ", "Nick"],
 ["Church rate", "Is the church promo price public (e.g., $4.99/mo)? If yes we show it on /churches/.", "churches page", "Nick"],
 ["Logo & key art", "Need: Kingdomland wordmark SVG, series thumbnails (16:10), character art (Lucy, Leo, RJ, Rylee), app screenshots for the device mockups.", "header/footer, show cards, devices", "Doris"],
 ["Legal", "Privacy Policy + Terms text (currently placeholders, noindex).", "/privacy/, /terms/", "Nick"],
 ["Hosting", "Where does the new site deploy (Netlify/Vercel/Cloudflare)? Who owns DNS for kingdomlandkids.com?", "—", "Khai"],
 ["Old URLs", "Export of current site URLs (or Screaming Frog crawl) for the redirect map.", "07_Redirect_Map", "Khai"],
]
sheet("08_Open_Questions", "Open questions · Câu hỏi mở trước go-live", "Answer these and the site is launch-ready. Each row names the file to edit.",
      ["Topic", "Question", "Where it lands", "Ask"], Q, [16, 90, 40, 12], wrap_cols=(2, 3))

wb.save(OUT)
print("wrote", OUT, "sheets:", wb.sheetnames)
